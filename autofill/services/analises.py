"""Leitura do FO.GSM.LMI.03.08 - Análises Microbiológicas de Água Potável."""

from __future__ import annotations

import io
import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
from openpyxl.utils import column_index_from_string

from autofill.model.config import Config, carregar_config, normalizar
from autofill.model.dados import Amostra, chave_ponto

LINHAS_CABECALHO = 15
CHAVES_OBRIGATORIAS = ("data", "origem", "identificacao")


class PlanilhaInvalida(ValueError):
    """O arquivo não tem a estrutura do FO.GSM.LMI.03.08."""


@dataclass
class Analises:
    aba: str
    amostras: list[Amostra]
    colunas: dict[str, str]  # chave -> letra usada
    ignoradas: list[int] = field(default_factory=list)  # linhas com dado mas sem data legível

    def do_ponto(self, origem: str, identificacao: str) -> list[Amostra]:
        chave = chave_ponto(origem, identificacao)
        return [a for a in self.amostras if a.chave == chave]

    @property
    def ultima_data(self) -> date | None:
        return max((a.data for a in self.amostras), default=None)


def _primeira_linha(valor: object) -> str:
    return normalizar(str(valor).splitlines()[0]) if valor is not None and str(valor).strip() else ""


def _mapear_colunas(linhas: list[tuple], config: Config) -> tuple[dict[str, int], int]:
    """Índice (base 0) de cada coluna pelo texto do cabeçalho; volta também a última linha do cabeçalho."""
    achadas: dict[str, int] = {}
    ultima = 0
    for n, linha in enumerate(linhas, start=1):
        for i, valor in enumerate(linha):
            texto = _primeira_linha(valor)
            if not texto:
                continue
            for chave, coluna in config.colunas.items():
                if chave not in achadas and texto in coluna.cabecalhos:
                    achadas[chave] = i
                    ultima = max(ultima, n)
    for chave, coluna in config.colunas.items():
        achadas.setdefault(chave, column_index_from_string(coluna.letra) - 1)
    return achadas, ultima


def _data(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        m = re.fullmatch(r"\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})\s*", valor)
        digitos = re.sub(r"\D", "", valor)
        if m is None and len(digitos) == 8:  # barra esquecida: "0510/2026" -> 05/10/2026
            m = re.fullmatch(r"(\d{2})(\d{2})(\d{4})", digitos)
        if m:
            dia, mes, ano = (int(g) for g in m.groups())
            try:
                return date(ano + 2000 if ano < 100 else ano, mes, dia)
            except ValueError:
                return None
    return None


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _escolher_aba(wb, config: Config):
    visiveis = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    preferida = [ws for ws in visiveis if normalizar(ws.title) == normalizar(config.aba_analises)]
    return (preferida or visiveis or wb.worksheets)[0]


def ler_analises(conteudo: bytes, config: Config | None = None) -> Analises:
    config = config or carregar_config()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception as erro:  # zip corrompido, formato antigo .xls etc.
            raise PlanilhaInvalida(f"Não foi possível abrir a planilha de análises: {erro}") from erro
    try:
        ws = _escolher_aba(wb, config)
        max_col = max(column_index_from_string(c.letra) for c in config.colunas.values()) + 20
        cabecalho = list(ws.iter_rows(min_row=1, max_row=LINHAS_CABECALHO, max_col=max_col, values_only=True))
        colunas, ultima = _mapear_colunas(cabecalho, config)
        if ultima == 0:
            raise PlanilhaInvalida(
                f"A aba '{ws.title}' não tem o cabeçalho do FO.GSM.LMI.03.08 (DATA, Origem, Identificação...)."
            )
        amostras: list[Amostra] = []
        ignoradas: list[int] = []
        for n, linha in enumerate(ws.iter_rows(min_row=ultima + 1, max_col=max_col, values_only=True), start=ultima + 1):
            valor = {chave: (linha[i] if i < len(linha) else None) for chave, i in colunas.items()}
            if valor["origem"] is None and valor["identificacao"] is None:
                continue
            data = _data(valor["data"])
            if data is None:
                ignoradas.append(n)
                continue
            amostras.append(
                Amostra(
                    linha=n,
                    data=data,
                    origem=_texto(valor["origem"]),
                    identificacao=_texto(valor["identificacao"]),
                    area=_texto(valor.get("area")),
                    setor=_texto(valor.get("setor")),
                    resultados={p.id: valor.get(p.id) for p in config.parametros},
                    rubrica=_texto(valor.get("rubrica")),
                )
            )
    finally:
        wb.close()
    letras = {chave: openpyxl.utils.get_column_letter(i + 1) for chave, i in colunas.items()}
    return Analises(ws.title, amostras, letras, ignoradas)

import io
import re
import zipfile
from datetime import date, datetime
from xml.dom import minidom

import openpyxl
import pytest

from autofill.model.config import carregar_config
from autofill.services import certificado as svc
from autofill.services.analises import PlanilhaInvalida, ler_analises
from autofill.services.exportacao import desempacotar, gerar_lote, zipar
from autofill.services.xlsx import PastaXlsx
from tests.conftest import COLETAS, planilha_analises

config = carregar_config()


@pytest.fixture
def analises(analises_bytes):
    return ler_analises(analises_bytes)


@pytest.fixture
def arquivo_417(certificado_417):
    return svc.ler_arquivo(certificado_417, "Vivência 417.xlsx")


def _abrir(xlsx: bytes):
    return openpyxl.load_workbook(io.BytesIO(xlsx))


# --------------------------------------------------------------------------- leitura das análises
def test_colunas_achadas_pelo_cabecalho(analises):
    assert analises.colunas["cloro"] == "BA"
    assert analises.colunas["ecoli"] == "BM"
    assert analises.aba == "Análises"


def test_data_sem_barra_e_recuperada_e_linha_sem_data_e_ignorada(analises):
    (a634,) = analises.do_ponto("Área de vivência", "634")
    assert a634.data == date(2026, 10, 5)
    assert analises.ignoradas == [19]


def test_ponto_numerico_e_texto_sao_o_mesmo(analises):
    assert len(analises.do_ponto("área de vivencia", 417)) == 3


def test_planilha_sem_cabecalho_e_recusada():
    wb = openpyxl.Workbook()
    wb.active["A1"] = "qualquer coisa"
    saida = io.BytesIO()
    wb.save(saida)
    with pytest.raises(PlanilhaInvalida):
        ler_analises(saida.getvalue())


# --------------------------------------------------------------------------- regras
@pytest.mark.parametrize(
    "parametro,valor,esperado",
    [
        ("coliformes", "Ausente", "Aceitável"),
        ("coliformes", "Presente", "Não aceitável"),
        ("cloro", 0.1, "Não aceitável"),
        ("cloro", 5.0, "Aceitável"),
        ("cloro", 5.1, "Não aceitável"),
        ("turbidez", "<1", "Aceitável"),
        ("cor", 16, "Não aceitável"),
        ("ph", 7.3, "Aceitável"),
        ("ph", None, "Conferir"),
        ("cloro", "ilegível", "Conferir"),
    ],
)
def test_avaliacao(parametro, valor, esperado):
    assert svc.avaliar(config.parametro(parametro), valor) == esperado


def test_abreviacoes_e_virgula_decimal():
    assert svc.valor_para_certificado(config.parametro("coliformes"), "A") == "Ausente"
    assert svc.valor_para_certificado(config.parametro("cloro"), "0,5") == 0.5
    assert svc.valor_para_certificado(config.parametro("cor"), 2.0) == 2


def test_validade_de_seis_meses_respeita_fim_do_mes():
    assert svc.somar_meses(date(2026, 10, 7), 6) == date(2027, 4, 7)
    assert svc.somar_meses(date(2026, 8, 31), 6) == date(2027, 2, 28)


def test_lista_de_pontos():
    assert svc.ler_lista_pontos("631, 634;665\n631") == ["631", "634", "665"]
    assert svc.ler_lista_pontos("600-602") == ["600", "601", "602"]


# --------------------------------------------------------------------------- certificado existente
def test_le_o_certificado_existente(arquivo_417):
    assert (arquivo_417.origem, arquivo_417.identificacao) == ("Área de vivência", "417")
    ultimo = arquivo_417.ultimo
    assert (ultimo.mes, ultimo.ano, ultimo.numero) == ("Outubro", 2026, 2)
    assert ultimo.linhas_parametros == {"coliformes": 23, "ecoli": 25, "cloro": 27, "turbidez": 29, "ph": 31, "cor": 33}


def test_mes_novo_vira_aba_nova_com_numero_seguinte(arquivo_417, analises):
    cert = svc.planejar(arquivo_417, analises, 9, 2026)
    assert (cert.aba_destino, cert.numero, cert.sobrescreve) == ("Setembro", 3, False)
    assert cert.amostra.data == date(2026, 9, 3)
    assert cert.validade == date(2027, 3, 3)
    assert cert.coletor == "Joel Borges - 18857"


def test_mesmo_mes_e_ano_atualiza_a_aba(arquivo_417, analises):
    cert = svc.planejar(arquivo_417, analises, 10, 2026)
    assert (cert.aba_destino, cert.numero, cert.sobrescreve) == ("Outubro", 2, True)


def test_mesmo_mes_de_outro_ano_ganha_o_ano_no_nome(arquivo_417, analises):
    cert = svc.planejar(arquivo_417, analises, 10, 2027)
    assert cert.aba_destino == "Outubro 2027"
    assert not cert.pronto


def test_recoleta_usa_a_coleta_mais_recente_do_mes(analises):
    cert = svc.planejar(svc.novo_arquivo("631"), analises, 10, 2026)
    assert len(cert.candidatas) == 2
    assert cert.amostra.data == date(2026, 10, 9)
    assert cert.aceitavel
    svc.escolher_amostra(cert, cert.candidatas[0])
    assert not cert.aceitavel
    assert any("Não aceitável" in a or "não aceitável" in a for a in cert.avisos)


def test_sem_coleta_pode_usar_a_anterior(arquivo_417, analises):
    assert not svc.planejar(arquivo_417, analises, 11, 2026).pronto
    cert = svc.planejar(arquivo_417, analises, 11, 2026, usar_anterior=True)
    assert cert.amostra.data == date(2026, 10, 7)


# --------------------------------------------------------------------------- geração do Excel
def test_gera_aba_nova_preservando_desenhos_e_assinatura(arquivo_417, analises):
    xlsx = svc.gerar_xlsx(svc.planejar(arquivo_417, analises, 9, 2026))
    partes = zipfile.ZipFile(io.BytesIO(xlsx)).namelist()
    assert len([p for p in partes if re.fullmatch(r"xl/drawings/drawing\d+\.xml", p)]) == 2
    assert "xl/media/image2.emf" in partes  # assinatura: o openpyxl a perderia
    for parte in partes:
        if parte.endswith((".xml", ".rels")):
            minidom.parseString(zipfile.ZipFile(io.BytesIO(xlsx)).read(parte))

    wb = _abrir(xlsx)
    assert wb.sheetnames[:2] == ["Outubro", "Setembro"]
    assert wb.active.title == "Setembro"
    ws = wb["Setembro"]
    assert (ws["M6"].value, ws["X6"].value, ws["AK6"].value) == ("Setembro", 2026, 3)
    assert ws["O14"].value.date() == date(2026, 9, 3)
    assert ws["O16"].value.date() == date(2027, 3, 3)
    assert ws["O14"].number_format == "dd/mm/yyyy"
    assert [ws[f"Y{n}"].value for n in (23, 25, 27, 29, 31, 33)] == ["Ausente", "Ausente", 1.8, 1, 7.5, 2]
    assert ws["AU27"].value == "Aceitável"
    assert wb["Outubro"]["Y27"].value == 1.6  # aba anterior intacta


def test_ponto_novo_parte_do_modelo(analises):
    cert = svc.planejar(svc.novo_arquivo("634"), analises, 10, 2026)
    wb = _abrir(svc.gerar_xlsx(cert))
    ws = wb["Outubro"]
    assert ws["O12"].value == "Área de vivência - 634"
    assert ws["AK6"].value == 1
    assert ws["Y27"].value == 0.1
    assert ws["AU27"].value == "Não aceitável"
    assert cert.arquivo.nome_arquivo == "Vivência 634.xlsx"


def test_lote_zipa_so_os_prontos(arquivo_417, analises):
    certs = [svc.planejar(arquivo_417, analises, 10, 2026)] + [
        svc.planejar(svc.novo_arquivo(p), analises, 10, 2026) for p in ("631", "999")
    ]
    planilhas, pdfs = gerar_lote(certs)
    assert set(planilhas) == {"Vivência 417.xlsx", "Vivência 631.xlsx"}
    assert pdfs == {}
    assert sorted(zipfile.ZipFile(io.BytesIO(zipar(planilhas))).namelist()) == sorted(planilhas)


def test_preparar_impressao_deixa_uma_aba_em_uma_pagina(certificado_417):
    pasta = PastaXlsx(certificado_417)
    pasta.preparar_impressao("Outubro")
    wb_xml = pasta._ler("xl/workbook.xml")
    assert "'Outubro'!$A$1:$BP$48" in wb_xml  # inclui a assinatura, que passa da coluna BE
    folha = pasta._ler(pasta.aba("Outubro").caminho)
    assert '<pageSetUpPr fitToPage="1"/>' in folha
    assert all(a.oculta for a in pasta.abas() if a.nome != "Outubro")


# --------------------------------------------------------------------------- .zip do mês anterior
def _zip(arquivos: dict[str, bytes]) -> bytes:
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w") as z:
        for nome, dados in arquivos.items():
            z.writestr(nome, dados)
    return saida.getvalue()


def test_planilha_solta_passa_direto(certificado_417):
    assert desempacotar("Vivência 417.xlsx", certificado_417) == ([("Vivência 417.xlsx", certificado_417)], [])


def test_zip_do_mes_anterior_vira_as_planilhas_de_dentro(certificado_417):
    pacote = _zip({
        "Vivência 417.xlsx": certificado_417,
        "pasta/Vivência 631.xlsm": b"631",
        "PDF/Vivência 417 - Outubro.pdf": b"%PDF",
        "~$Vivência 417.xlsx": b"temporario",
        "__MACOSX/._Vivência 417.xlsx": b"lixo",
    })
    planilhas, avisos = desempacotar("Certificados_Outubro_2026.zip", pacote)
    assert [n for n, _ in planilhas] == ["Vivência 417.xlsx", "Vivência 631.xlsm"]
    assert planilhas[0][1] == certificado_417
    assert avisos == []


def test_zip_sem_planilha_ou_corrompido_avisa():
    assert desempacotar("vazio.zip", _zip({"PDF/a.pdf": b"%PDF"}))[1] == ["vazio.zip: nenhuma planilha .xlsx dentro do .zip."]
    assert desempacotar("ruim.zip", b"nao e zip")[0] == []


def test_ciclo_mensal_com_o_zip_baixado(arquivo_417, analises):
    """Outubro gera o .zip; em Novembro esse mesmo .zip entra e cada planilha ganha a aba nova."""
    coletas_novembro = [(datetime(2026, 11, 5), "Área de vivência", 417, 0, 0.3, 7.4, 1.1, "Ausente", "Ausente")]
    analises_novembro = ler_analises(planilha_analises(COLETAS + coletas_novembro))
    planilhas, _ = gerar_lote([svc.planejar(arquivo_417, analises, 10, 2026)])
    (nome, conteudo), = desempacotar("Certificados_Outubro_2026.zip", zipar(planilhas))[0]

    cert = svc.planejar(svc.ler_arquivo(conteudo, nome), analises_novembro, 11, 2026)
    wb = _abrir(svc.gerar_xlsx(cert))
    assert wb.sheetnames[:2] == ["Outubro", "Novembro"]
    assert (wb["Novembro"]["AK6"].value, wb["Novembro"]["Y27"].value) == (3, 1.1)

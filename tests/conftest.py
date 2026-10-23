"""Planilha de análises sintética no layout do FO.GSM.LMI.03.08 (cabeçalho nas linhas 8 a 10)."""

from __future__ import annotations

import io
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from autofill.model.config import CAMINHO_MODELO  # noqa: E402

# (data, origem, identificação, cor, turbidez, pH, cloro, coliformes, termotolerantes)
COLETAS = [
    (datetime(2026, 4, 14), "Área de vivência", 417, 1, 0.2, 7.5, 1.4, "Ausente", "Ausente"),
    (datetime(2026, 9, 3), "Área de vivência", 417, 2, 1, 7.5, 1.8, "Ausente", "Ausente"),
    (datetime(2026, 10, 7), "Área de vivência", 417, 1, 0.4, 7.3, 1.6, "Ausente", "Ausente"),
    (datetime(2026, 10, 2), "Área de vivência", 631, 0, 0.3, 7.0, 1.0, "Presente", "Presente"),
    (datetime(2026, 10, 9), "Área de vivência", 631, 0, 0.3, 7.1, 1.2, "Ausente", "Ausente"),
    ("0510/2026", "Área de vivência", 634, 3, 0.5, 7.2, 0.1, "Ausente", "Ausente"),
    (datetime(2026, 10, 6), "Caminhão pipa", 665, 0, 0.2, 6.9, 0.9, "A", "A"),
    (None, "Equipe", 30, 0, 0.2, 6.9, 0.9, "Ausente", "Ausente"),
]


def planilha_analises(coletas=COLETAS) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Análises"
    ws["H1"] = "FORMULÁRIO"
    ws["H3"] = "Data de Publicação:"
    ws["A8"], ws["G8"], ws["AK8"] = "DATA", "AMOSTRA", "ENSAIOS FÍSICO QUÍMICO"
    ws["AK9"], ws["AQ9"], ws["AV9"], ws["BA9"] = "Cor Aparente", "Turbidez", "pH", " Cloro"
    ws["BV8"] = "Rubrica"
    ws["G10"], ws["N10"], ws["AD10"] = "Origem", "Identificação", "Setor"
    ws["BF10"] = "Coliformes Totais\nAusente \n(em 100 ml)"
    ws["BM10"] = "Coliformes Termotolerantes \nAusente \n(em 100 ml)"
    for n, (data, origem, ident, cor, turb, ph, cloro, col, ecoli) in enumerate(coletas, start=12):
        for letra, valor in zip(
            ("A", "G", "N", "V", "AD", "AK", "AQ", "AV", "BA", "BF", "BM", "BV"),
            (data, origem, ident, "TRATOS CULTURAIS", "PÁTIO CTA", cor, turb, ph, cloro, col, ecoli, "Paulo/Evelyn"),
        ):
            ws[f"{letra}{n}"] = valor
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


@pytest.fixture
def analises_bytes() -> bytes:
    return planilha_analises()


@pytest.fixture
def certificado_417() -> bytes:
    """O modelo é um certificado real da Vivência 417 (aba Outubro/2026, nº 2)."""
    return CAMINHO_MODELO.read_bytes()
